import os
import json
import time
import shutil
import threading
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

from flask import Flask, request, jsonify, send_file, send_from_directory

# ─── Config ───────────────────────────────────────────────────────────────────
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'livox2026')
GITHUB_TOKEN   = os.environ.get('GITHUB_TOKEN', '')
GITHUB_REPO    = os.environ.get('GITHUB_REPO', '')

BASE   = Path(__file__).parent
DATA   = BASE / 'data.json'
OPLOG  = BASE / 'oplog.json'
BACKUP = BASE / 'backup'
BACKUP.mkdir(exist_ok=True)

app = Flask(__name__, static_folder=None)

# ─── Data helpers ─────────────────────────────────────────────────────────────

def read_data():
    if DATA.exists():
        try:
            return json.loads(DATA.read_text(encoding='utf-8'))
        except Exception:
            return []
    return []


def write_data(rows):
    # backup
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    shutil.copy2(DATA, BACKUP / f'data_{ts}.json') if DATA.exists() else None
    DATA.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')
    if GITHUB_TOKEN and GITHUB_REPO:
        threading.Thread(target=_sync_to_github, daemon=True).start()


def read_log():
    if OPLOG.exists():
        try:
            return json.loads(OPLOG.read_text(encoding='utf-8'))
        except Exception:
            return []
    return []


def append_log(entry):
    logs = read_log()
    logs.insert(0, entry)
    logs = logs[:500]
    OPLOG.write_text(json.dumps(logs, ensure_ascii=False, indent=2), encoding='utf-8')


def _sync_to_github():
    try:
        import urllib.request, base64
        url = f'https://api.github.com/repos/{GITHUB_REPO}/contents/data.json'
        headers = {
            'Authorization': f'token {GITHUB_TOKEN}',
            'Content-Type': 'application/json',
            'User-Agent': 'kanban-sync',
        }
        # Get current SHA
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=10) as r:
            info = json.loads(r.read())
        sha = info.get('sha', '')
        # PUT update
        content_b64 = base64.b64encode(DATA.read_bytes()).decode()
        payload = json.dumps({
            'message': f'auto-sync {datetime.now().strftime("%Y-%m-%d %H:%M")}',
            'content': content_b64,
            'sha': sha,
        }).encode()
        req2 = urllib.request.Request(url, data=payload, headers=headers, method='PUT')
        urllib.request.urlopen(req2, timeout=15)
    except Exception:
        pass


def get_caller(req, body=None):
    raw = req.headers.get('X-Person', '')
    if not raw and body:
        raw = body.get('caller', '') or ''
    try:
        return unquote(raw).strip() or '匿名'
    except Exception:
        return str(raw).strip() or '匿名'


def is_admin(req):
    token = req.headers.get('X-Admin-Token', '')
    return token == ADMIN_PASSWORD


def next_id(rows):
    if not rows:
        return 1
    return max(r.get('id', 0) for r in rows) + 1


def now_str():
    return datetime.now().strftime('%Y-%m-%d %H:%M')

# ─── Routes ───────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    return send_from_directory(BASE, 'kanban.html')


@app.route('/health')
def health():
    return jsonify({'status': 'ok'})


@app.route('/api/data')
def api_data():
    return jsonify(read_data())


@app.route('/api/auth', methods=['POST'])
def api_auth():
    body = request.get_json(silent=True) or {}
    pwd = body.get('password', '')
    if pwd == ADMIN_PASSWORD:
        return jsonify({'ok': True, 'token': ADMIN_PASSWORD})
    return jsonify({'ok': False, 'msg': '密码错误'}), 401


@app.route('/api/log')
def api_log():
    if not is_admin(request):
        return jsonify({'error': '无权限'}), 403
    q = request.args.get('q', '').strip()
    logs = read_log()
    if q:
        logs = [l for l in logs if q in json.dumps(l, ensure_ascii=False)]
    return jsonify(logs)


@app.route('/api/public_update/<int:item_id>', methods=['POST'])
def api_public_update(item_id):
    body = request.get_json(silent=True) or {}
    caller = get_caller(request, body)
    rows = read_data()
    target = next((r for r in rows if r['id'] == item_id), None)
    if not target:
        return jsonify({'error': '未找到'}), 404

    allowed = {'person', 'progress', 'submit_url', 'status', 'livox_confirm'}
    changed = {}
    for k in allowed:
        if k in body:
            old = target.get(k, '')
            target[k] = body[k]
            if str(old) != str(body[k]):
                changed[k] = {'old': old, 'new': body[k]}

    target['updated_at'] = now_str()
    target['updated_by'] = caller
    write_data(rows)
    append_log({
        'time': now_str(),
        'op': '公开填写',
        'item': target.get('item', ''),
        'caller': caller,
        'changed': changed,
    })
    return jsonify({'ok': True})


@app.route('/api/update/<int:item_id>', methods=['POST'])
def api_update(item_id):
    body = request.get_json(silent=True) or {}
    caller = get_caller(request, body)
    rows = read_data()
    target = next((r for r in rows if r['id'] == item_id), None)
    if not target:
        return jsonify({'error': '未找到'}), 404

    editable = {'date', 'item', 'submit', 'ddl', 'person', 'livox',
                'progress', 'submit_url', 'status', 'livox_confirm'}
    changed = {}
    for k in editable:
        if k in body:
            old = target.get(k, '')
            target[k] = body[k]
            if str(old) != str(body[k]):
                changed[k] = {'old': old, 'new': body[k]}

    target['updated_at'] = now_str()
    target['updated_by'] = caller
    write_data(rows)
    append_log({
        'time': now_str(),
        'op': '编辑事项',
        'item': target.get('item', ''),
        'caller': caller,
        'changed': changed,
    })
    return jsonify({'ok': True})


@app.route('/api/add', methods=['POST'])
def api_add():
    if not is_admin(request):
        return jsonify({'error': '无权限'}), 403
    body = request.get_json(silent=True) or {}
    caller = get_caller(request, body)
    rows = read_data()
    new_row = {
        'id': next_id(rows),
        'date': body.get('date', now_str()[:10]),
        'item': body.get('item', ''),
        'submit': body.get('submit', ''),
        'ddl': body.get('ddl', ''),
        'person': body.get('person', ''),
        'livox': body.get('livox', ''),
        'progress': body.get('progress', ''),
        'submit_url': body.get('submit_url', ''),
        'status': body.get('status', '未完成'),
        'livox_confirm': body.get('livox_confirm', '未完成'),
        'updated_at': now_str(),
        'updated_by': caller,
    }
    rows.append(new_row)
    write_data(rows)
    append_log({
        'time': now_str(),
        'op': '新增事项',
        'item': new_row['item'],
        'caller': caller,
        'changed': {},
    })
    return jsonify({'ok': True, 'id': new_row['id']})


@app.route('/api/delete/<int:item_id>', methods=['DELETE'])
def api_delete(item_id):
    if not is_admin(request):
        return jsonify({'error': '无权限'}), 403
    caller = get_caller(request)
    rows = read_data()
    target = next((r for r in rows if r['id'] == item_id), None)
    if not target:
        return jsonify({'error': '未找到'}), 404
    rows = [r for r in rows if r['id'] != item_id]
    write_data(rows)
    append_log({
        'time': now_str(),
        'op': '删除事项',
        'item': target.get('item', ''),
        'caller': caller,
        'changed': {},
    })
    return jsonify({'ok': True})


@app.route('/api/template')
def api_template():
    tpl = BASE / 'template.xlsx'
    if tpl.exists():
        return send_file(tpl, as_attachment=True, download_name='template.xlsx')
    # Generate minimal template on the fly
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '事项模板'
        headers = ['序号', '日期(YYYY-MM-DD)', '事项', '提交物需求', 'DDL', '责任人', 'LIVOX对接人', '进展/闭环情况', '提交物链接', '状态(未完成/完成/挂起)']
        ws.append(headers)
        ws.append(['(从第2行开始填写数据)'])
        tmp = BASE / 'template.xlsx'
        wb.save(tmp)
        return send_file(tmp, as_attachment=True, download_name='template.xlsx')
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/batch_add', methods=['POST'])
def api_batch_add():
    if not is_admin(request):
        return jsonify({'error': '无权限'}), 403
    body = request.get_json(silent=True) or {}
    caller = get_caller(request, body)
    parse_only = body.get('parse_only', False)

    items = body.get('items')
    if items is None:
        # Try base64 Excel
        b64 = body.get('excel_b64', '')
        if not b64:
            return jsonify({'error': '无数据'}), 400
        try:
            import base64, io, openpyxl
            raw = base64.b64decode(b64)
            wb = openpyxl.load_workbook(io.BytesIO(raw))
            ws = wb.active
            items = []
            for r in range(2, ws.max_row + 1):
                vals = [ws.cell(r, c).value for c in range(1, 11)]
                if not any(v is not None for v in vals):
                    continue
                ddl_v = vals[4]
                date_v = vals[1]
                items.append({
                    'date': date_v.strftime('%Y-%m-%d') if isinstance(date_v, datetime) else str(date_v or ''),
                    'item': str(vals[2] or ''),
                    'submit': str(vals[3] or ''),
                    'ddl': ddl_v.strftime('%Y-%m-%d') if isinstance(ddl_v, datetime) else str(ddl_v or ''),
                    'person': str(vals[5] or ''),
                    'livox': str(vals[6] or ''),
                    'progress': str(vals[7] or ''),
                    'submit_url': str(vals[8] or ''),
                    'status': str(vals[9] or '未完成'),
                    'livox_confirm': '未完成',
                })
        except Exception as e:
            return jsonify({'error': f'Excel解析失败: {e}'}), 400

    if parse_only:
        return jsonify({'ok': True, 'preview': items, 'count': len(items)})

    rows = read_data()
    added = []
    for it in items:
        new_row = {
            'id': next_id(rows),
            'date': it.get('date', ''),
            'item': it.get('item', ''),
            'submit': it.get('submit', ''),
            'ddl': it.get('ddl', ''),
            'person': it.get('person', ''),
            'livox': it.get('livox', ''),
            'progress': it.get('progress', ''),
            'submit_url': it.get('submit_url', ''),
            'status': it.get('status', '未完成'),
            'livox_confirm': it.get('livox_confirm', '未完成'),
            'updated_at': now_str(),
            'updated_by': caller,
        }
        rows.append(new_row)
        added.append(new_row)

    write_data(rows)
    append_log({
        'time': now_str(),
        'op': '批量新增',
        'item': f'共{len(added)}条',
        'caller': caller,
        'changed': {},
    })
    return jsonify({'ok': True, 'count': len(added)})


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=8080, debug=False)
